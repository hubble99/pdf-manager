"""Builds the PDF-to-Word sample PDFs.

Run with: python make_fixtures.py  (needs reportlab and Pillow)
Each PDF covers one layout case. The Rust tests convert them and check that
the expected text, tables and images survive; README.md records how each one
looks in Word.
"""
import io
import os
import sys

from PIL import Image, ImageDraw
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image as RLImage, ListFlowable, ListItem, PageBreak,
    PageTemplate, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

OUT = os.path.dirname(os.path.abspath(__file__)) if len(sys.argv) < 2 else sys.argv[1]
styles = getSampleStyleSheet()
body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=14, alignment=TA_JUSTIFY)

LOREM = (
    "Dokuwu converts this paragraph from PDF to Word. The sentence continues long enough to wrap "
    "over several lines, so the converter must join the lines back into one paragraph instead of "
    "breaking it at every line end. Numbers like 1,234.56 and symbols such as % & @ must survive."
)


def picture(w, h, hue):
    img = Image.new("RGB", (w, h), hue)
    d = ImageDraw.Draw(img)
    for i in range(0, w, 20):
        d.line([(i, 0), (w - i, h)], fill=(255, 255, 255), width=2)
    d.ellipse([w // 4, h // 4, 3 * w // 4, 3 * h // 4], outline=(0, 0, 0), width=4)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def text_pdf():
    doc = SimpleDocTemplate(os.path.join(OUT, "text.pdf"), pagesize=A4, title="Text sample")
    s = [Paragraph("Quarterly Report", styles["Title"]), Paragraph("1. Introduction", styles["Heading1"])]
    s += [Paragraph(LOREM, body), Spacer(1, 6)]
    s += [Paragraph("This line has <b>bold words</b>, <i>italic words</i> and <u>underlined words</u>.", body)]
    s += [Paragraph("1.1 Key points", styles["Heading2"])]
    s += [ListFlowable([ListItem(Paragraph(x, body)) for x in
                        ["First bullet point", "Second bullet point with more text", "Third bullet point"]],
                       bulletType="bullet")]
    s += [PageBreak(), Paragraph("2. Second page", styles["Heading1"])]
    s += [Paragraph(LOREM, body) for _ in range(3)]
    doc.build(s)


def columns_pdf():
    path = os.path.join(OUT, "columns.pdf")
    doc = BaseDocTemplate(path, pagesize=A4, title="Two columns")
    w, h = A4
    m = 2 * cm
    gap = 0.8 * cm
    colw = (w - 2 * m - gap) / 2
    frames = [Frame(m, m, colw, h - 2 * m, id="left"), Frame(m + colw + gap, m, colw, h - 2 * m, id="right")]
    doc.addPageTemplates([PageTemplate(id="two", frames=frames)])
    s = [Paragraph("Two-column article", styles["Heading1"])]
    for i in range(1, 9):
        s.append(Paragraph(f"Section {i}", styles["Heading3"]))
        s.append(Paragraph(f"COLUMN-PARA-{i}. " + LOREM, body))
    doc.build(s)


def table_pdf():
    doc = SimpleDocTemplate(os.path.join(OUT, "table.pdf"), pagesize=A4, title="Tables")
    data = [["Item", "Qty", "Unit price", "Total"]]
    for i in range(1, 9):
        data.append([f"Product {i}", str(i), f"{i * 1.5:.2f}", f"{i * i * 1.5:.2f}"])
    data.append(["Grand total", "", "", "306.00"])
    t = Table(data, colWidths=[6 * cm, 2 * cm, 3 * cm, 3 * cm])
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe7ff")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("SPAN", (0, -1), (2, -1)),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
    ]))
    # A table drawn without any lines: only alignment says it is a table.
    plain = Table([["Name", "Role", "City"], ["Ana", "Editor", "Bandung"], ["Budi", "Designer", "Surabaya"]],
                  colWidths=[4 * cm, 4 * cm, 4 * cm])
    s = [Paragraph("Invoice table (with grid lines)", styles["Heading2"]), t, Spacer(1, 20),
         Paragraph("Borderless table", styles["Heading2"]), plain]
    doc.build(s)


def images_pdf():
    doc = SimpleDocTemplate(os.path.join(OUT, "images.pdf"), pagesize=A4, title="Images")
    s = [Paragraph("Images", styles["Heading1"]), Paragraph("A photo-like image follows.", body)]
    s += [RLImage(picture(600, 400, (40, 110, 200)), width=12 * cm, height=8 * cm),
          Paragraph("Figure 1: blue picture.", styles["Italic"]), Spacer(1, 12)]
    s += [RLImage(picture(300, 300, (200, 80, 60)), width=5 * cm, height=5 * cm),
          Paragraph("Figure 2: red picture.", styles["Italic"])]
    doc.build(s)


def mixed_pdf():
    path = os.path.join(OUT, "mixed.pdf")

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, A4[1] - 1.2 * cm, "ACME Corp · Internal")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=A4, title="Mixed")
    data = [["Metric", "Q1", "Q2"], ["Revenue", "120", "140"], ["Cost", "80", "85"]]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
    s = [Paragraph("Mixed layout", styles["Title"]), Paragraph(LOREM, body),
         RLImage(picture(400, 200, (60, 160, 90)), width=10 * cm, height=5 * cm),
         Paragraph("Results", styles["Heading2"]), t, Paragraph(LOREM, body), PageBreak(),
         Paragraph("Appendix", styles["Heading1"]), Paragraph(LOREM, body)]
    doc.build(s, onFirstPage=decorate, onLaterPages=decorate)


def scanned_pdf():
    # A page that is only a picture of text: no text layer, so nothing to extract without OCR.
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    for i in range(20):
        d.text((100, 100 + i * 60), "Scanned line of text " + str(i + 1), fill="black")
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=70)
    buf.seek(0)
    doc = SimpleDocTemplate(os.path.join(OUT, "scanned.pdf"), pagesize=A4, title="Scanned",
                            leftMargin=0, rightMargin=0, topMargin=0, bottomMargin=0)
    doc.build([RLImage(buf, width=A4[0] - 40, height=A4[1] - 40)])


def protected_pdf():
    # Needs a password to open; the converter must say so instead of failing oddly.
    from reportlab.lib import pdfencrypt
    enc = pdfencrypt.StandardEncryption("secret", canPrint=1)
    doc = SimpleDocTemplate(os.path.join(OUT, "protected.pdf"), pagesize=A4, encrypt=enc)
    doc.build([Paragraph("Secret text.", body)])


def word_export_pdf():
    """A document as Word would export it: embedded subset fonts, a header,
    a numbered list, a table with merged cells and shading, and a photo.
    Needs python-docx and LibreOffice (soffice); skipped when missing."""
    import shutil
    import subprocess
    import tempfile
    try:
        import docx
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        print("skipping word-export.pdf: python-docx missing")
        return
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("skipping word-export.pdf: LibreOffice missing")
        return
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Dokuwu sample · Confidential"
    d.add_heading("Project Proposal", 0)
    d.add_heading("Background", 1)
    p = d.add_paragraph(LOREM + " ")
    r = p.add_run("A red, bold sentence follows the plain text.")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x20, 0x20)
    d.add_paragraph(LOREM)
    d.add_heading("Goals", 2)
    for item in ["Reduce manual work", "Keep every document editable", "Ship before the end of the quarter"]:
        d.add_paragraph(item, style="List Number")
    d.add_heading("Budget", 1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate([["Item", "Q1", "Q2"], ["Design", "1,200", "800"], ["Build", "3,400", "2,100"], ["Total", "", "7,500"]]):
        for j, val in enumerate(row):
            t.cell(i, j).text = val
    t.cell(3, 0).merge(t.cell(3, 1))
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "D9E2F3")
    t.cell(0, 0)._tc.get_or_add_tcPr().append(shading)
    d.add_paragraph("")
    # A photo-like image: smooth gradients compress as JPEG.
    img = Image.new("RGB", (480, 300))
    px = img.load()
    for x in range(480):
        for y in range(300):
            px[x, y] = ((x * 255) // 480, (y * 255) // 300, ((x + y) * 3) % 256)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    buf.seek(0)
    d.add_picture(buf, width=Cm(10))
    cap = d.add_paragraph("Figure 1: gradient photo.")
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    d.add_paragraph(LOREM)
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "word-export.docx")
        d.save(src)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                       check=True, capture_output=True, timeout=180)
        shutil.copy(os.path.join(tmp, "word-export.pdf"), os.path.join(OUT, "word-export.pdf"))


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for fn in (text_pdf, columns_pdf, table_pdf, images_pdf, mixed_pdf, scanned_pdf, protected_pdf, word_export_pdf):
        fn()
    print("written to", OUT)
